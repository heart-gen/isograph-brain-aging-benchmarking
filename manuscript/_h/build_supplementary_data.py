"""Export the manuscript supplementary tables and Supplementary Data files.

Nature-family and Cell-family journals keep only small tables inside the
Supplementary Information PDF. Everything larger is submitted as separate
machine-readable files (cited here as "Data S1", "Data S2", ...),
one Excel workbook per item, each with a README sheet that carries the item's
title, legend and provenance.

This script maps every repo-numbered CSV (written by ``assemble_supp_tables.py``
or by ``isograph_benchmark/figures/synthetic_benchmark.R``) to its manuscript
item and writes a manuscript-ready export tree::

    manuscript/_m/manuscript_supplement/
        supplementary_tables/tableS01_<slug>.csv           # stay in the PDF
        supplementary_data/dataS01_<slug>.xlsx                # one workbook per item
        supplementary_data/dataS01_<slug>.csv                 # same table as CSV
        MANIFEST.md                                        # item -> source mapping

The two directories are copied verbatim into the manuscript repository
(``content/supplementary_tables/`` and ``content/supplementary_data/``).

Legends for the README sheets are read from the manuscript supplement
(``content/99.supplement.md``) when it is available, so the workbook text is
identical to the published legend. Run from the repo root::

    python manuscript/_h/build_supplementary_data.py \
        [--supplement ../../manuscripts/isograph-brain-manuscript/content/99.supplement.md]
"""
from __future__ import annotations

import argparse
import datetime as dt
import re
import shutil
from dataclasses import dataclass
from pathlib import Path

import pandas as pd
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

REPO = Path(__file__).resolve().parents[2]
SUPP_TABLES = REPO / "manuscript" / "_m" / "supp_tables"
BENCH_TABLES = REPO / "01_synthetic_benchmark" / "03_metrics" / "_m"
OUT = REPO / "manuscript" / "_m" / "manuscript_supplement"
REPO_URL = "https://github.com/heart-gen/isograph-brain-aging-benchmarking"

# Tables with at most 40 rows and 12 columns stay in the PDF as Supplementary
# Tables; everything else becomes a Supplementary Data file.
MAX_ROWS, MAX_COLS = 40, 12


@dataclass(frozen=True)
class Item:
    kind: str        # "table" (in the PDF) or "data" (separate file)
    number: int      # manuscript number within its kind
    ref_id: str      # pandoc identifier used in the manuscript (tbl:<ref_id>)
    slug: str        # file-name stem
    source: Path     # repo CSV
    title: str

    @property
    def label(self) -> str:
        return f"Table S{self.number}" if self.kind == "table" else f"Data S{self.number}"

    @property
    def stem(self) -> str:
        if self.kind == "table":
            return f"tableS{self.number:02d}_{self.slug}"
        return f"dataS{self.number:02d}_{self.slug}"

    @property
    def subdir(self) -> str:
        return "supplementary_tables" if self.kind == "table" else "supplementary_data"


def _t(ref: str, slug: str, src: Path, title: str) -> tuple:
    return ("table", ref, slug, src, title)


def _d(ref: str, slug: str, src: Path, title: str) -> tuple:
    return ("data", ref, slug, src, title)


def _number(entries: list[tuple]) -> list[Item]:
    counters = {"table": 0, "data": 0}
    items = []
    for kind, ref, slug, src, title in entries:
        counters[kind] += 1
        items.append(Item(kind, counters[kind], ref, slug, src, title))
    return items


ITEMS: list[Item] = _number([
    # Entries are in order of first citation (Results, Discussion, Methods,
    # Supplement); numbers are assigned from that order within each kind.
    # ---- Supplementary Tables (rendered in the PDF) --------------------------
    _t("s-benchmark", "benchmark_summary", BENCH_TABLES / "tableS_benchmark_summary.csv",
       "Synthetic benchmark performance across scenarios and methods."),
    _t("s-scale-compute", "scale_compute_summary", BENCH_TABLES / "tableS_scale_compute_summary.csv",
       "Runtime and memory use across network sizes and compute settings."),
    _t("s-separation-axis", "axis_orthogonality", SUPP_TABLES / "tableS13a_axis_orthogonality.csv",
       "Separation of the switch and abundance coordinates, per analysis."),
    _t("s-baseline-pooled", "baseline_pooled", SUPP_TABLES / "tableS1_baseline_pooled.csv",
       "Pooled per-module association rates across the 17 applications."),
    _t("s-clinical", "clinical_consequence", SUPP_TABLES / "tableS37_clinical_consequence.csv",
       "Constraint and ClinVar density of switch genes and switched exons."),
    _t("s-qtl-contrast", "qtl_specificity_contrast", SUPP_TABLES / "tableS3_qtl_specificity_contrast.csv",
       "Splicing-QTL specificity contrast, all analyses."),
    _t("s-qtl-baseline", "qtl_specificity_matched_baseline", SUPP_TABLES / "tableS4_qtl_specificity_matched_baseline.csv",
       "Splicing-QTL specificity contrast on the analyses shared by all methods."),
    _t("s-module-cis", "module_cis_control", SUPP_TABLES / "tableS23_module_cis_control.csv",
       "Module-level cis control of isoform choice."),
    _t("s-synthetic", "synthetic_scenarios", SUPP_TABLES / "tableS35_synthetic_scenarios.csv",
       "Synthetic benchmark scenarios."),
    # ---- Supplementary Data (separate files) ---------------------------------
    _d("s-cohorts", "cohort_description", SUPP_TABLES / "tableS34_cohort_description.csv",
       "Discovery cohorts."),
    _d("s-composition-adj", "composition_adjustment", SUPP_TABLES / "tableS13_composition_adjustment.csv",
       "Switch-unique genes before and after adjustment for estimated cellular composition."),
    _d("s-trust-funnel", "module_trust_funnel", SUPP_TABLES / "tableS7_module_trust_funnel.csv",
       "Per-region module trust funnel."),
    _d("s-baseline-region", "baseline_per_region", SUPP_TABLES / "tableS2_baseline_per_region.csv",
       "Per-analysis module association rates for all four module sets."),
    _d("s-split-half", "split_half_module_ledger", SUPP_TABLES / "tableS7a_split_half_module_ledger.csv",
       "Split-half module ledger."),
    _d("s-resolution", "resolution_sensitivity", SUPP_TABLES / "tableS7e_resolution_sensitivity.csv",
       "Split-half agreement across the Leiden resolution sweep."),
    _d("s-projection", "projection_module_ledger", SUPP_TABLES / "tableS7b_projection_module_ledger.csv",
       "Cross-cohort eigengene projection ledger."),
    _d("s-crosscohort-perm", "crosscohort_permutation", SUPP_TABLES / "tableS7c_crosscohort_permutation.csv",
       "Cross-cohort matched-pair count against its permutation nulls."),
    _d("s-functional", "functional_preservation", SUPP_TABLES / "tableS7d_functional_preservation.csv",
       "Functional preservation of matched cross-cohort module pairs."),
    _d("s-saturn", "saturn_concordance", SUPP_TABLES / "tableS15_isa_concordance.csv",
       "satuRn concordance of IsoGraph switch genes, per analysis."),
    _d("s-psi-genes", "psi_gene_corroboration", SUPP_TABLES / "tableS38_psi_gene_corroboration.csv",
       "Gene-level corroboration of the switch layer by junction usage."),
    _d("s-heldout-dtu", "module_context_heldout_dtu", SUPP_TABLES / "tableS21_module_context_heldout_dtu.csv",
       "Module context and held-out DTU evidence, per analysis."),
    _d("s-rbp-regulons", "rbp_regulons", SUPP_TABLES / "tableS32_rbp_regulons.csv",
       "RNA-binding-protein motif regulons of the switch modules."),
    _d("s-rbp-eclip", "rbp_eclip_binding", SUPP_TABLES / "tableS33_rbp_eclip_binding.csv",
       "ENCODE eCLIP binding at switched and constitutive exons."),
    _d("s-allelic", "allelic_imbalance_regions", SUPP_TABLES / "tableS22_allelic_imbalance_regions.csv",
       "Within-donor allelic test of isoform choice, per region."),
    _d("s-signal-coloc", "signal_coloc_nominations", SUPP_TABLES / "tableS25_signal_coloc_nominations.csv",
       "Signal-level sQTL colocalization nominations."),
    _d("s-coloc-contrast", "coloc_modality_contrast", SUPP_TABLES / "tableS26_coloc_modality_contrast.csv",
       "Paired sQTL and eQTL colocalization contrast."),
    _d("s-brainseq-coloc", "brainseq_axis_coloc_contrast", SUPP_TABLES / "tableS31_brainseq_axis_coloc_contrast.csv",
       "BrainSEQ switch- and abundance-axis colocalization contrast."),
    _d("s-coloc-events", "coloc_isoform_events", SUPP_TABLES / "tableS28_coloc_isoform_events.csv",
       "Colocalization events and their mapping to IsoGraph switch pairs."),
    _d("s-smr", "smr_heidi", SUPP_TABLES / "tableS30_smr_heidi.csv",
       "SMR and HEIDI for the colocalization nominations."),
    _d("s-clpp-events", "clpp_isoform_events", SUPP_TABLES / "tableS24_clpp_isoform_events.csv",
       "CLPP-nominated isoform events and their mapping to IsoGraph switch pairs."),
    _d("s-ldsc", "ldsc_partitioned", SUPP_TABLES / "tableS27_ldsc_partitioned.csv",
       "Partitioned heritability of the switch-derived QTL annotations."),
    _d("s-coloc-convergence", "coloc_convergence", SUPP_TABLES / "tableS20a_coloc_convergence_global.csv",
       "Module concentration of colocalizing switch genes."),
    _d("s-magma", "magma_module_gwas", SUPP_TABLES / "tableS36_magma_module_gwas.csv",
       "MAGMA competitive gene-set tests of every module."),
    _d("s-longread", "longread_coloc_confirmation", SUPP_TABLES / "tableS29_longread_coloc_confirmation.csv",
       "Long-read confirmation of colocalization-anchored switch pairs."),
    _d("s-junction-pairs", "junction_pair_corroboration", SUPP_TABLES / "tableS40_junction_pair_corroboration.csv",
       "Short-read junction corroboration of colocalization-prioritized transcript pairs."),
    _d("s-sign-scale", "projection_sign_scale", SUPP_TABLES / "tableS7f_projection_sign_scale.csv",
       "Raw against null-standardized projected-age sign agreement, per region."),
    _d("s-psi-junctions", "psi_junction_confirmation", SUPP_TABLES / "tableS39_psi_junction_confirmation.csv",
       "Short-read junction confirmation of the CLPP-anchored switch pairs."),
])


# --------------------------------------------------------------------------
# Legends from the manuscript supplement
# --------------------------------------------------------------------------
def _figure_numbers(supplement: Path) -> dict[str, str]:
    """Map pandoc figure ids to their printed numbers (main text + supplement)."""
    numbers: dict[str, str] = {}
    results = supplement.with_name("03.results.md")
    if results.exists():
        for i, m in enumerate(re.finditer(r"\{#fig:([\w-]+)", results.read_text()), start=1):
            numbers.setdefault(m.group(1), str(i))
    for m in re.finditer(r'\{#fig:([\w-]+)\s+tag="(S\d+)"', supplement.read_text()):
        numbers[m.group(1)] = m.group(2)
    return numbers


def _plain(text: str, fig_numbers: dict[str, str], tbl_numbers: dict[str, str],
           data_numbers: dict[str, str]) -> str:
    """Reduce manuscript Markdown to plain text for a README sheet."""
    text = re.sub(r"\{@fig:([\w-]+)\}", lambda m: fig_numbers.get(m.group(1), "?"), text)
    text = re.sub(r"\{@tbl:([\w-]+)\}", lambda m: tbl_numbers.get(m.group(1), "?"), text)
    text = re.sub(r"\{@data:([\w-]+)\}", lambda m: data_numbers.get(m.group(1), "?"), text)
    text = re.sub(r"\[([^\]]+)\]\([^)]+\)", r"\1", text)      # links -> text
    text = text.replace("**", "")
    text = re.sub(r"(?<!\w)\*(\S[^*]*?)\*(?!\w)", r"\1", text)  # *italic*
    text = re.sub(r"\^([^^]+)\^", r"\1", text)                  # ^superscript^
    text = text.replace("\\|", "|")
    return re.sub(r"\s+", " ", text).strip()


def read_legends(supplement: Path | None) -> dict[str, str]:
    """Return {ref_id: plain legend} parsed from the manuscript supplement."""
    if supplement is None or not supplement.exists():
        return {}
    text = supplement.read_text()
    fig_numbers = _figure_numbers(supplement)
    tbl_numbers = {m.group(1): m.group(2)
                   for m in re.finditer(r'\{#tbl:([\w-]+)\s+tag="(S\d+)"', text)}
    # Data items are numbered by build/pandoc/filters/datanos.lua in document order
    data_numbers = {m.group(1): f"S{i}"
                    for i, m in enumerate(re.finditer(r"^::: \{#data:([\w-]+)\}", text, re.M), start=1)}
    legends: dict[str, str] = {}
    # Supplementary Tables: "Table: **Title.** legend {#tbl:id tag="Sn"}"
    for m in re.finditer(r"^Table: \*\*(.+?)\*\*\s*(.*?)\s*\{#tbl:([\w-]+)[^}]*\}", text, re.S | re.M):
        legends[m.group(3)] = _plain(m.group(2), fig_numbers, tbl_numbers, data_numbers)
    # Supplementary Data: "::: {#data:id}\n**Title.** legend\n\n[Download ...]\n:::"
    for m in re.finditer(r"^::: \{#data:([\w-]+)\}\n\*\*(.+?)\*\*\s*(.*?)(?=\n\n|\n:::)", text, re.S | re.M):
        legends[m.group(1)] = _plain(m.group(3), fig_numbers, tbl_numbers, data_numbers)
    return legends


def check_order(supplement: Path | None) -> None:
    """Warn when the manuscript numbering would not match ITEMS.

    Supplementary items are numbered by their order of definition in the
    supplement; Nature-family journals expect that order to follow the first
    citation in the text (Results, Discussion, Methods, then the supplement).
    """
    if supplement is None or not supplement.exists():
        return
    text = supplement.read_text()
    defined = {"table": [m.group(1) for m in re.finditer(r'\{#tbl:([\w-]+)\s+tag=', text)],
               "data": [m.group(1) for m in re.finditer(r"^::: \{#data:([\w-]+)\}", text, re.M)]}
    for kind, refs in defined.items():
        expected = [it.ref_id for it in ITEMS if it.kind == kind]
        if refs != expected:
            print(f"WARNING: {kind} items are defined in the supplement in a different order than ITEMS:\n"
                  f"  supplement: {refs}\n  ITEMS:      {expected}")
    cited: list[str] = []
    for name in ["03.results.md", "04.discussion.md", "05.methods.md", "99.supplement.md"]:
        p = supplement.with_name(name)
        if p.exists():
            for m in re.finditer(r"\{@(?:tbl|data):([\w-]+)\}", p.read_text()):
                if m.group(1) not in cited:
                    cited.append(m.group(1))
    for kind in ("table", "data"):
        expected = [it.ref_id for it in ITEMS if it.kind == kind]
        by_citation = [r for r in cited if r in expected]
        if by_citation != expected:
            print(f"WARNING: {kind} items are not numbered in order of first citation:\n"
                  f"  first citation: {by_citation}\n  ITEMS:          {expected}")


# --------------------------------------------------------------------------
# Writers
# --------------------------------------------------------------------------
def write_workbook(item: Item, df: pd.DataFrame, legend: str, path: Path) -> None:
    generated = dt.date.today().isoformat()
    readme = pd.DataFrame(
        [
            ("Item", item.label),
            ("Title", item.title),
            ("Legend", legend or "(see Supplementary Information)"),
            ("Rows", len(df)),
            ("Columns", df.shape[1]),
            ("Column names", "; ".join(map(str, df.columns))),
            ("Source file", str(item.source.relative_to(REPO))),
            ("Analysis repository", REPO_URL),
            ("Generated", f"{generated} by manuscript/_h/build_supplementary_data.py"),
        ],
        columns=["Field", "Value"],
    )
    with pd.ExcelWriter(path, engine="openpyxl") as xl:
        readme.to_excel(xl, sheet_name="README", index=False)
        df.to_excel(xl, sheet_name="Data", index=False)
        ws = xl.sheets["README"]
        ws.column_dimensions["A"].width = 22
        ws.column_dimensions["B"].width = 110
        for row in ws.iter_rows(min_row=2, max_col=2):
            row[1].alignment = Alignment(wrap_text=True, vertical="top")
            row[0].alignment = Alignment(vertical="top")
            row[0].font = Font(bold=True)
        ws["A1"].font = ws["B1"].font = Font(bold=True)
        ws = xl.sheets["Data"]
        ws.freeze_panes = "A2"
        for j, col in enumerate(df.columns, start=1):
            ws.cell(row=1, column=j).font = Font(bold=True)
            sample = df[col].astype(str).head(200).map(len).max() if len(df) else 0
            ws.column_dimensions[get_column_letter(j)].width = min(60, max(10, int(max(len(str(col)), sample) * 1.1) + 2))


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument(
        "--supplement",
        type=Path,
        default=REPO.parent.parent / "manuscripts" / "isograph-brain-manuscript" / "content" / "99.supplement.md",
        help="manuscript content/99.supplement.md, used for the README legends",
    )
    args = parser.parse_args()

    legends = read_legends(args.supplement)
    check_order(args.supplement)
    if OUT.exists():
        shutil.rmtree(OUT)
    (OUT / "supplementary_tables").mkdir(parents=True)
    (OUT / "supplementary_data").mkdir(parents=True)

    manifest = ["# Manuscript supplementary items", "",
                f"Written by `manuscript/_h/build_supplementary_data.py` on {dt.date.today().isoformat()}.",
                f"Supplementary Tables stay in the PDF (<= {MAX_ROWS} rows and <= {MAX_COLS} columns); "
                "Supplementary Data are separate Excel workbooks (README + Data sheets) with CSV copies.", "",
                "| Item | Manuscript file | Source CSV | Rows | Columns | Title |", "| --- | --- | --- | --- | --- | --- |"]
    for item in ITEMS:
        if not item.source.exists():
            raise SystemExit(f"missing source for {item.label}: {item.source}")
        df = pd.read_csv(item.source)
        if item.kind == "table" and (len(df) > MAX_ROWS or df.shape[1] > MAX_COLS):
            raise SystemExit(f"{item.label} is {df.shape} and no longer fits the PDF; move it to Supplementary Data")
        dest_dir = OUT / item.subdir
        shutil.copyfile(item.source, dest_dir / f"{item.stem}.csv")
        if item.kind == "data":
            write_workbook(item, df, legends.get(item.ref_id, ""), dest_dir / f"{item.stem}.xlsx")
        manifest.append(f"| {item.label} | `{item.subdir}/{item.stem}"
                        f"{'.xlsx' if item.kind == 'data' else '.csv'}` | "
                        f"`{item.source.relative_to(REPO)}` | {len(df)} | {df.shape[1]} | {item.title} |")
        print(f"{item.label:22} {item.stem:55} {len(df):6} x {df.shape[1]:2}"
              f"{'' if legends.get(item.ref_id) else '  (no legend found)'}")
    (OUT / "MANIFEST.md").write_text("\n".join(manifest) + "\n")
    print(f"\nwrote {len(ITEMS)} items to {OUT.relative_to(REPO)}")


if __name__ == "__main__":
    main()
